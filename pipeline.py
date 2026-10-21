
import json
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
import numpy as np
from moviepy import ImageClip, concatenate_videoclips


def extract_content(file_path):
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext in [".txt", ".md"]:
        return path.read_text(encoding="utf-8", errors="ignore")

    if ext == ".pdf":
        import fitz
        doc = fitz.open(file_path)
        return "\n\n".join(page.get_text() for page in doc)

    if ext == ".docx":
        from docx import Document
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(file_path)
        result = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    result.append(shape.text)
        return "\n".join(result)

    if ext in [".png", ".jpg", ".jpeg", ".webp"]:
        try:
            import pytesseract
            return pytesseract.image_to_string(
                Image.open(file_path)
            )
        except Exception:
            return "Image uploaded. Topic description is needed."

    raise ValueError("Unsupported file format")


def generate_storyboard(
    content, topic, duration_minutes, language, style, add_labels=True
):
    from openai import OpenAI

    client = OpenAI(api_key=os.environ["OPENAI_API_KEY"])

    prompt = f"""
Create a medical educational video storyboard.
Topic: {topic or 'Infer from source'}
Source: {content[:20000]}
Duration: {duration_minutes} minutes
Language: {language}
Visual style: {style}
Include labels: {add_labels}

Return JSON only:
{{"scenes":[
  {{"title":"Introduction",
    "visual_description":"What viewers see",
    "narration":"Educational narration",
    "labels":["Anatomical structure"],
    "duration_seconds":30}}
]}}

Create enough scenes to cover the topic clearly.
Explain normal anatomy, mechanism, body effects and recap.
Do not invent facts unsupported by the source.
"""

    response = client.chat.completions.create(
        model=os.getenv("OPENAI_MODEL", "gpt-4.1-mini"),
        messages=[
            {"role": "system", "content": "Return valid JSON only."},
            {"role": "user", "content": prompt},
        ],
        response_format={"type": "json_object"},
    )

    data = json.loads(response.choices[0].message.content)
    scenes = data.get("scenes", [])

    if not scenes:
        raise ValueError("No scenes generated")

    total = sum(max(1, int(s.get("duration_seconds", 10))) for s in scenes)

    for scene in scenes:
        old = max(1, int(scene.get("duration_seconds", 10)))
        scene["duration_seconds"] = max(
            1, round(old * duration_minutes * 60 / total)
        )
        scene.setdefault("title", "Medical scene")
        scene.setdefault("visual_description", "")
        scene.setdefault("narration", "")
        scene.setdefault("labels", [])

    return scenes


def render_video(
    scenes, output_path, duration_minutes, resolution="720p",
    disclaimer=True
):
    width, height = (
        (1920, 1080) if resolution == "1080p" else (1280, 720)
    )
    font = ImageFont.load_default()
    clips = []

    durations = [
        max(1, int(s.get("duration_seconds", 10))) for s in scenes
    ]
    total = sum(durations)
    durations = [
        max(1, round(d * duration_minutes * 60 / total))
        for d in durations
    ]
    durations[-1] += duration_minutes * 60 - sum(durations)

    try:
        for i, scene in enumerate(scenes):
            image = Image.new("RGB", (width, height), (5, 18, 32))
            draw = ImageDraw.Draw(image)

            draw.rectangle((0, 0, width, 8), fill=(0, 210, 220))
            draw.text(
                (60, 50),
                f"MEDCINE 3D STUDIO | SCENE {i + 1}",
                font=font,
                fill=(0, 220, 220),
            )
            draw.text(
                (60, 120), scene["title"][:100],
                font=font, fill="white"
            )

            y = 190
            for text in [
                scene["visual_description"][:350],
                "Labels: " + ", ".join(scene.get("labels", []))[:300],
                "Narration: " + scene["narration"][:350],
            ]:
                draw.text((60, y), text, font=font, fill=(200, 220, 235))
                y += 90

            if disclaimer:
                draw.text(
                    (60, height - 40),
                    "Educational visualization - verify medical content",
                    font=font, fill=(100, 200, 200),
                )

            clips.append(
                ImageClip(np.array(image)).with_duration(durations[i])
            )

        final = concatenate_videoclips(clips, method="compose")
        final.write_videofile(
            output_path, fps=24, codec="libx264",
            audio=False, logger=None
        )
        final.close()

    finally:
        for clip in clips:
            try:
                clip.close()
            except Exception:
                pass
